from __future__ import annotations

import os
import re
import hashlib
import random
import shutil
import uuid
from io import BytesIO
from functools import lru_cache
from pathlib import Path
from typing import Any

import pymupdf
import psycopg
from fastapi import Depends, FastAPI, File, HTTPException, Request, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, Response, StreamingResponse
from fastapi.security import HTTPAuthorizationCredentials, HTTPBearer
from pydantic import BaseModel, Field
import mimetypes
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.util import Inches, Pt

from auth import create_token, decode_token, hash_password, is_valid_email, token_secret, verify_password
from chunking import chunk_text
from database import (create_user, find_user_by_email, get_paper, get_user, hybrid_search,
                      create_chat_session, get_user_chat, initialise, list_admin_logins, list_admin_users, list_papers, list_user_chats,
                      list_paper_annotations, list_user_data, list_user_notifications, get_user_settings,
                      mark_user_notifications_read, update_user_settings, record_activity, record_login, save_artifact,
                      save_paper_annotation, delete_paper_annotation,
                      save_chat_message, save_chunks, save_paper_if_unique, sync_admin_emails,
                      find_duplicate_paper, delete_user_collection, list_user_collections,
                      save_user_collection, update_user_collection, user_owns_artifacts, list_paper_chunks,
                      update_flashcard_reviews, start_study_session, add_study_session_time,
                      save_quiz_attempt, get_user_artifact, get_dashboard_metrics)
from rag import (ai_provider_status, answer, citations_for, embed, embed_query,
                 expand_mindmap_node, generate_study_artifact)
from document_processor import extract_text

BASE_DIR = Path(__file__).resolve().parent
PAPERS_DIR = Path(os.getenv("PAPER_STORAGE_DIR", str(BASE_DIR / "data" / "papers"))).resolve()
MAX_UPLOAD_BYTES = 100 * 1024 * 1024
ALLOWED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".pptx",
    ".txt",
    ".png",
    ".jpg",
    ".jpeg",
}
AUTH = HTTPBearer(auto_error=False)
ARTIFACT_TYPES = {"flashcards", "mindmap", "quiz", "summary", "report", "ppt_outline", "viva", "literature_review", "visualization", "comparison", "research_gap", "research_ideas", "reading_guide", "method_extraction"}

app = FastAPI(title="ScholarMind Backend", version="1.1.0")
app.add_middleware(CORSMiddleware, allow_origins=os.getenv("CORS_ORIGINS", "http://localhost:5173").split(","),
                   allow_credentials=True, allow_methods=["*"], allow_headers=["*"])


class RegisterRequest(BaseModel):
    name: str = Field(min_length=2, max_length=120)
    email: str = Field(min_length=5, max_length=254)
    password: str = Field(min_length=8, max_length=256)
    role: str = Field(default="student", pattern="^(student|researcher|professor|industry)$")


class LoginRequest(BaseModel):
    email: str = Field(min_length=5, max_length=254)
    password: str = Field(min_length=1, max_length=256)


class ChatRequest(BaseModel):
    question: str = Field(min_length=2, max_length=4000)
    paper_ids: list[str] | None = None
    session_id: str | None = None
    top_k: int = Field(default=5, ge=1, le=10)


class GenerateRequest(BaseModel):
    kind: str = Field(pattern="^(flashcards|mindmap|quiz|summary|report|ppt_outline|viva|literature_review|visualization|comparison|research_gap|research_ideas|reading_guide|method_extraction)$")
    paper_id: str | None = None
    paper_ids: list[str] | None = Field(default=None, max_length=10)
    prompt: str = Field(default="", max_length=1000)
    count: int = Field(default=6, ge=1, le=20)
    difficulty: str = Field(default="medium", pattern="^(simple|medium|hard)$")
    avoid_questions: list[str] = Field(default_factory=list, max_length=20)
    variation_seed: str | None = Field(default=None, max_length=64)
    presentation_input_type: str = Field(default="paper", pattern="^(paper|notes|topic)$")
    source_text: str = Field(default="", max_length=15000)
    presentation_purpose: str = Field(default="research", pattern="^(research|lecture|thesis_defense|conference|project_pitch)$")


class MindMapExistingNode(BaseModel):
    id: str = Field(min_length=1, max_length=120)
    label: str = Field(min_length=1, max_length=180)


class MindMapExpansionRequest(BaseModel):
    paper_id: str
    node_id: str = Field(min_length=1, max_length=120)
    node_label: str = Field(min_length=2, max_length=180)
    breadcrumb: list[str] = Field(default_factory=list, max_length=12)
    context: str = Field(default="", max_length=1600)
    existing_nodes: list[MindMapExistingNode] = Field(default_factory=list, max_length=250)


class MindMapExpansionNode(BaseModel):
    id: str = Field(min_length=1, max_length=120)
    label: str = Field(min_length=2, max_length=100)
    summary: str = Field(min_length=20, max_length=500)
    details: str = Field(min_length=30, max_length=1200)
    example: str = Field(default="", max_length=500)
    key_points: list[str] = Field(default_factory=list, max_length=4)
    related_concepts: list[str] = Field(default_factory=list, max_length=4)
    importance: bool = False
    page: int = Field(ge=1)
    evidence: str = Field(min_length=20, max_length=180)


class MindMapExpansionEdge(BaseModel):
    source: str = Field(min_length=1, max_length=120)
    target: str = Field(min_length=1, max_length=120)
    label: str = Field(min_length=2, max_length=80)


class MindMapExpansionResponse(BaseModel):
    parent_id: str
    children: list[MindMapExpansionNode]
    edges: list[MindMapExpansionEdge]
    generation_mode: str = "ai"


class ArtifactUpdate(BaseModel):
    title: str | None = Field(default=None, max_length=200)
    payload: dict[str, Any] | None = None


class PresentationSlide(BaseModel):
    title: str = Field(min_length=1, max_length=200)
    content: list[str] = Field(default_factory=list, max_length=30)
    speaker_notes: str = Field(default="", max_length=10000)
    page: int | None = Field(default=None, ge=1)
    pages: list[int] = Field(default_factory=list, max_length=30)


class PresentationExportRequest(BaseModel):
    title: str = Field(default="ScholarMind Presentation", min_length=1, max_length=200)
    source_label: str = Field(default="Research presentation", max_length=100)
    slides: list[PresentationSlide] = Field(min_length=1, max_length=30)


class FlashcardReviewRequest(BaseModel):
    reviews: dict[str, str] = Field(min_length=1, max_length=20)


class StudySessionHeartbeatRequest(BaseModel):
    elapsed_seconds: int = Field(ge=0, le=60)
    end_session: bool = False


class QuizAttemptRequest(BaseModel):
    answers: dict[str, str] = Field(default_factory=dict, max_length=20)


class PaperAnnotationRequest(BaseModel):
    page_number: int = Field(ge=1)
    kind: str = Field(pattern="^(highlight|note)$")
    content: str = Field(min_length=1, max_length=5000)


class SettingsUpdate(BaseModel):
    profile: dict[str, str] | None = None
    notifications: dict[str, bool] | None = None
    theme: str | None = Field(default=None, pattern="^(light|dark|system)$")


class ResearchFeedbackRequest(BaseModel):
    category: str = Field(pattern="^(answer_relevance|citation_support|evidence_quality)$")
    rating: str = Field(pattern="^(helpful|needs_revision|direct|indirect|supported|partially_supported|unclear)$")
    session_id: str | None = None
    paper_id: str | None = None
    page_number: int | None = Field(default=None, ge=1)
    note: str = Field(default="", max_length=500)


class ResearchCollectionRequest(BaseModel):
    name: str = Field(min_length=1, max_length=100)
    paper_ids: list[str] = Field(default_factory=list, max_length=50)
    artifact_ids: list[str] = Field(default_factory=list, max_length=100)


@app.on_event("startup")
def startup() -> None:
    PAPERS_DIR.mkdir(parents=True, exist_ok=True)
    initialise()
    sync_admin_emails()


def current_user(credentials: HTTPAuthorizationCredentials | None = Depends(AUTH)) -> dict[str, Any]:
    if not credentials:
        raise HTTPException(status_code=401, detail="Sign in to continue")
    try:
        claims = decode_token(credentials.credentials)
        user = get_user(claims["sub"])
    except (ValueError, RuntimeError):
        user = None
    if not user:
        raise HTTPException(status_code=401, detail="Your session has expired. Sign in again.")
    return user


def require_auth_configuration() -> None:
    """Fail cleanly before touching account data if token signing is misconfigured."""
    try:
        token_secret()
    except RuntimeError as exc:
        raise HTTPException(
            status_code=503,
            detail="Sign-in is temporarily unavailable because server authentication is not configured.",
        ) from exc


def admin_user(user: dict[str, Any] = Depends(current_user)) -> dict[str, Any]:
    if user["role"] != "admin":
        raise HTTPException(status_code=403, detail="Administrator access is required")
    return user


def owned_paper(paper_id: str, user: dict[str, Any]) -> dict[str, Any]:
    paper = get_paper(paper_id)
    if not paper or paper.get("owner_id") != user["id"]:
        raise HTTPException(status_code=404, detail="Paper not found")
    return paper


@lru_cache(maxsize=128)
def pdf_index(file_path: str) -> dict[str, Any]:
    """Return real PDF bookmarks and a page index with headings detected from page text."""
    with pymupdf.open(file_path) as document:
        page_count = len(document)
        detected: list[list[str]] = [[] for _ in range(page_count)]
        heading_pattern = re.compile(r"^\d+(?:\.\d+)*\.?\s+[A-Z][\w ,:;()&/'’–-]{1,100}$")
        named_pattern = re.compile(
            r"^(abstract|introduction|background|related work|method(?:s|ology)?|"
            r"experiments?|results?(?: and discussion)?|discussion|conclusion|"
            r"references|acknowledg(?:e)?ments|appendix(?:\s+[A-Z])?)$", re.IGNORECASE
        )

        for page_number, page in enumerate(document, start=1):
            page_index = page_number - 1
            for block in page.get_text("dict").get("blocks", []):
                for line in block.get("lines", []):
                    spans = line.get("spans", [])
                    title = re.sub(r"\s+", " ", "".join(span.get("text", "") for span in spans)).strip()
                    if not title or len(title) > 110:
                        continue
                    bold = any(span.get("flags", 0) & 16 for span in spans)
                    size = max((span.get("size", 0) for span in spans), default=0)
                    is_heading = bool(heading_pattern.match(title) or named_pattern.match(title))
                    is_styled_heading = bold and size >= 12 and len(title.split()) <= 12 and not title.endswith((".", ";", ","))
                    if (is_heading or is_styled_heading) and title not in detected[page_index]:
                        detected[page_index].append(title)

        toc = document.get_toc()
        contents = [
            {"title": title.strip(), "page": page_number, "level": max(level - 1, 0)}
            for level, title, page_number, *_ in toc
            if title.strip() and 1 <= page_number <= page_count
        ]
        if not contents:
            contents = [
                {"title": title, "page": page_number, "level": 0}
                for page_number, titles in enumerate(detected, start=1)
                for title in titles
            ]

        pages = [
            {"page": page_number, "title": titles[0] if titles else f"Page {page_number}"}
            for page_number, titles in enumerate(detected, start=1)
        ]
        return {"contents": contents, "pages": pages}


def sha256_file(file_path: Path) -> str:
    digest = hashlib.sha256()
    with file_path.open("rb") as source:
        for block in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def duplicate_upload_result(paper: dict[str, Any]) -> dict[str, Any]:
    return {"paper_id": paper["id"], "filename": paper["filename"],
            "total_pages": paper["page_count"],"file_type": paper.get("file_type"), "status": "duplicate",
            "duplicate_of": paper["id"]}


@app.get("/api/health")
def health_check():
    return {"status": "ok", "message": "ScholarMind backend is running"}


@app.get("/")
def root():
    return {
        "message": "ScholarMind backend is running",
        "docs": "/docs",
        "health": "/api/health",
    }


@app.get("/favicon.ico", include_in_schema=False)
def favicon():
    return Response(status_code=204)


@app.post("/api/auth/register", status_code=201)
def register(payload: RegisterRequest):
    require_auth_configuration()
    email = payload.email.strip().lower()
    if not is_valid_email(email):
        raise HTTPException(status_code=422, detail="Enter a valid email address")
    if find_user_by_email(email):
        raise HTTPException(status_code=409, detail="An account with this email already exists")
    admins = {value.strip().lower() for value in os.getenv("ADMIN_EMAILS", "").split(",") if value.strip()}
    role = "admin" if email in admins else payload.role
    try:
        user = create_user(str(uuid.uuid4()), payload.name.strip(), email, hash_password(payload.password), role)
    except psycopg.errors.UniqueViolation as exc:
        raise HTTPException(status_code=409, detail="An account with this email already exists") from exc
    return {"access_token": create_token(user), "token_type": "bearer", "user": user}


@app.post("/api/auth/login")
def login(payload: LoginRequest, request: Request):
    require_auth_configuration()
    email = payload.email.strip().lower()
    user = find_user_by_email(email)
    succeeded = bool(user and verify_password(payload.password, user["password_hash"]))
    record_login(email, succeeded, user["id"] if succeeded else None,
                 request.client.host if request.client else None, request.headers.get("user-agent"))
    if not succeeded:
        raise HTTPException(status_code=401, detail="Email or password is incorrect")
    public_user = {key: user[key] for key in ("id", "name", "email", "role", "created_at", "last_login_at")}
    return {"access_token": create_token(user), "token_type": "bearer", "user": public_user}


@app.get("/api/auth/me")
def me(user: dict[str, Any] = Depends(current_user)):
    return {"user": user}


@app.get("/api/me/data")
def my_data(user: dict[str, Any] = Depends(current_user)):
    return list_user_data(user["id"])


@app.get("/api/collections")
def collections(user: dict[str, Any] = Depends(current_user)):
    return {"collections": list_user_collections(user["id"])}


def persist_collection(payload: ResearchCollectionRequest, user: dict[str, Any], collection_id: str | None = None) -> dict[str, Any]:
    name = payload.name.strip()
    if not name:
        raise HTTPException(status_code=422, detail="Enter a collection name")
    paper_ids = list(dict.fromkeys(payload.paper_ids))
    artifact_ids = list(dict.fromkeys(payload.artifact_ids))
    try:
        for paper_id in paper_ids:
            uuid.UUID(paper_id)
        for item_id in artifact_ids:
            uuid.UUID(item_id)
    except ValueError as exc:
        raise HTTPException(status_code=422, detail="A saved-work ID is invalid") from exc
    for paper_id in paper_ids:
        owned_paper(paper_id, user)
    if not user_owns_artifacts(user["id"], artifact_ids):
        raise HTTPException(status_code=404, detail="One or more saved items were not found")
    result = (update_user_collection(user["id"], collection_id, name, paper_ids, artifact_ids)
              if collection_id else save_user_collection(user["id"], name, paper_ids, artifact_ids))
    if not result:
        raise HTTPException(status_code=404, detail="Collection not found")
    record_activity(user["id"], "saved_research_collection", {"collection_id": result["id"], "name": name})
    return result


@app.post("/api/collections", status_code=201)
def create_collection(payload: ResearchCollectionRequest, user: dict[str, Any] = Depends(current_user)):
    try:
        return {"collection": persist_collection(payload, user)}
    except psycopg.errors.UniqueViolation as exc:
        raise HTTPException(status_code=409, detail="A collection with that name already exists") from exc


@app.put("/api/collections/{collection_id}")
def update_collection(collection_id: str, payload: ResearchCollectionRequest, user: dict[str, Any] = Depends(current_user)):
    existing = next((item for item in list_user_collections(user["id"]) if item["id"] == collection_id), None)
    if not existing:
        raise HTTPException(status_code=404, detail="Collection not found")
    try:
        return {"collection": persist_collection(payload, user, collection_id)}
    except psycopg.errors.UniqueViolation as exc:
        raise HTTPException(status_code=409, detail="A collection with that name already exists") from exc


@app.delete("/api/collections/{collection_id}", status_code=204)
def remove_collection(collection_id: str, user: dict[str, Any] = Depends(current_user)):
    if not delete_user_collection(user["id"], collection_id):
        raise HTTPException(status_code=404, detail="Collection not found")
    return Response(status_code=204)

@app.post("/api/study-sessions", status_code=201)
def begin_study_session(user: dict[str, Any] = Depends(current_user)):
    return {"session_id": start_study_session(user["id"])}


@app.get("/api/dashboard/metrics")
def dashboard_metrics(user: dict[str, Any] = Depends(current_user)):
    return get_dashboard_metrics(user["id"])


@app.post("/api/study-sessions/{session_id}/heartbeat")
def study_session_heartbeat(session_id: str, payload: StudySessionHeartbeatRequest,
                            user: dict[str, Any] = Depends(current_user)):
    if not add_study_session_time(user["id"], session_id, payload.elapsed_seconds, payload.end_session):
        raise HTTPException(status_code=404, detail="Study session not found or already ended")
    return {"ok": True}


@app.get("/api/settings")
def settings(user: dict[str, Any] = Depends(current_user)):
    return get_user_settings(user["id"])


@app.patch("/api/settings")
def save_settings(payload: SettingsUpdate, user: dict[str, Any] = Depends(current_user)):
    allowed_profile = {"name", "institution", "bio", "field", "keywords", "scholar", "github", "linkedin"}
    allowed_notifications = {"papers", "ideas", "viva", "newsletter", "marketing"}
    profile = payload.profile or {}
    notifications = payload.notifications or {}
    if set(profile) - allowed_profile:
        raise HTTPException(status_code=422, detail="The profile contains unsupported fields")
    if set(notifications) - allowed_notifications:
        raise HTTPException(status_code=422, detail="The notification preferences contain unsupported fields")
    profile = {key: value.strip() for key, value in profile.items()}
    limits = {"name": 120, "institution": 200, "bio": 1000, "field": 120,
              "keywords": 500, "scholar": 500, "github": 500, "linkedin": 500}
    for key, value in profile.items():
        if len(value) > limits[key]:
            raise HTTPException(status_code=422, detail=f"{key.title()} is too long")
    if "name" in profile and len(profile["name"]) < 2:
        raise HTTPException(status_code=422, detail="Name must contain at least 2 characters")
    return update_user_settings(user["id"], profile=profile, notifications=notifications, theme=payload.theme)


@app.get("/api/notifications")
def notifications(user: dict[str, Any] = Depends(current_user)):
    items = list_user_notifications(user["id"])
    return {"notifications": items, "unread_count": sum(not item["read"] for item in items)}


@app.post("/api/notifications/read")
def read_notifications(user: dict[str, Any] = Depends(current_user)):
    mark_user_notifications_read(user["id"])
    return {"ok": True}


@app.get("/api/chats")
def chats(user: dict[str, Any] = Depends(current_user)):
    return {"chats": list_user_chats(user["id"])}


@app.get("/api/chats/{chat_id}")
def chat_detail(chat_id: str, user: dict[str, Any] = Depends(current_user)):
    result = get_user_chat(user["id"], chat_id)
    if not result:
        raise HTTPException(status_code=404, detail="Chat not found")
    return {"chat": result}


@app.post("/api/research-feedback", status_code=201)
def research_feedback(payload: ResearchFeedbackRequest, user: dict[str, Any] = Depends(current_user)):
    if payload.session_id and not get_user_chat(user["id"], payload.session_id):
        raise HTTPException(status_code=404, detail="Chat session not found")
    if payload.paper_id:
        owned_paper(payload.paper_id, user)
    details = {"category": payload.category, "rating": payload.rating,
               "session_id": payload.session_id, "paper_id": payload.paper_id,
               "page_number": payload.page_number, "note": payload.note.strip()}
    record_activity(user["id"], "research_feedback", details)
    return {"ok": True}


@app.get("/api/admin/overview")
def admin_overview(user: dict[str, Any] = Depends(admin_user)):
    return {"users": list_admin_users(), "login_activity": list_admin_logins()}


@app.get("/api/papers")
def papers(user: dict[str, Any] = Depends(current_user)):
    return {"papers": [paper for paper in list_papers() if paper.get("owner_id") == user["id"]]}


@app.get("/api/papers/{paper_id}")
def paper(paper_id: str, user: dict[str, Any] = Depends(current_user)):
    result = owned_paper(paper_id, user)
    return {key: value for key, value in result.items() if key != "stored_path"}


@app.get("/api/papers/{paper_id}/index")
def paper_index(paper_id: str, user: dict[str, Any] = Depends(current_user)):
    result = owned_paper(paper_id, user)
    if not Path(result["stored_path"]).is_file():
        raise HTTPException(status_code=404, detail="Paper file not found")
    try:
        if result["file_type"] != ".pdf":
            return {
             "contents": [],
            "pages": [],
            "message": "Indexing currently supported only for PDF files."
    }
        return pdf_index(result["stored_path"])
    except Exception as exc:
        raise HTTPException(status_code=422, detail=f"Could not read PDF contents: {exc}") from exc


@app.get("/api/papers/{paper_id}/annotations")
def paper_annotations(paper_id: str, user: dict[str, Any] = Depends(current_user)):
    owned_paper(paper_id, user)
    return {"annotations": list_paper_annotations(paper_id, user["id"])}


@app.post("/api/papers/{paper_id}/annotations", status_code=201)
def create_paper_annotation(paper_id: str, payload: PaperAnnotationRequest,
                            user: dict[str, Any] = Depends(current_user)):
    paper = owned_paper(paper_id, user)
    if payload.page_number > paper["page_count"]:
        raise HTTPException(status_code=422, detail="Page number is outside this PDF")
    content = payload.content.strip()
    if not content:
        raise HTTPException(status_code=422, detail="Enter a quote or note before saving")
    annotation = save_paper_annotation(str(uuid.uuid4()), paper_id, user["id"],
                                       payload.page_number, payload.kind, content)
    return {"annotation": annotation}


@app.delete("/api/papers/{paper_id}/annotations/{annotation_id}", status_code=204)
def remove_paper_annotation(paper_id: str, annotation_id: str,
                            user: dict[str, Any] = Depends(current_user)):
    owned_paper(paper_id, user)
    if not delete_paper_annotation(annotation_id, paper_id, user["id"]):
        raise HTTPException(status_code=404, detail="Annotation not found")
    return Response(status_code=204)


@app.get("/api/papers/{paper_id}/file")
def paper_file(paper_id: str, user: dict[str, Any] = Depends(current_user)):
    result = owned_paper(paper_id, user)

    if not Path(result["stored_path"]).is_file():
        raise HTTPException(status_code=404, detail="File not found")

    media_type, _ = mimetypes.guess_type(
        result["filename"]
    )

    return FileResponse(
        result["stored_path"],
        media_type=media_type or "application/octet-stream",
        filename=result["filename"]
    )


@app.post("/api/papers/upload", status_code=201)
def upload_paper(
    file: UploadFile = File(...),
    user: dict[str, Any] = Depends(current_user),
):
    if not file.filename:
        raise HTTPException(
            status_code=400,
            detail="No file was selected",
        )

    original_name = Path(file.filename).name
    extension = Path(original_name).suffix.lower()

    if extension not in ALLOWED_EXTENSIONS:
        raise HTTPException(
            status_code=415,
            detail=(
                "Unsupported file type. Allowed formats: "
                "PDF, DOCX, PPTX, TXT, PNG, JPG, JPEG."
            ),
        )

    paper_id = str(uuid.uuid4())

    safe_name = original_name

    file_path = PAPERS_DIR / f"{paper_id}{extension}"

    with file_path.open("wb") as destination:
        shutil.copyfileobj(file.file, destination)

    if file_path.stat().st_size > MAX_UPLOAD_BYTES:
        file_path.unlink(missing_ok=True)

        raise HTTPException(
            status_code=413,
            detail="File must be 100 MB or smaller",
        )

    try:
        # ---------------------------------------------------------
        # 1. Calculate file hash
        # ---------------------------------------------------------
        content_sha256 = sha256_file(file_path)

        # ---------------------------------------------------------
        # 2. Check duplicate
        # ---------------------------------------------------------
        duplicate = find_duplicate_paper(
            user["id"],
            content_sha256,
        )

        if duplicate:
            file_path.unlink(missing_ok=True)

            record_activity(
                user["id"],
                "duplicate_upload_skipped",
                {
                    "paper_id": duplicate["id"],
                    "filename": safe_name,
                },
            )

            return duplicate_upload_result(duplicate)

        # ---------------------------------------------------------
        # 3. Extract text
        # ---------------------------------------------------------
        extracted = extract_text(file_path)

        full_text = extracted["text"]
        pages = extracted["pages"]
        page_count = extracted["page_count"]

        if not full_text.strip():
            raise ValueError(
                "No readable text was found in this file."
            )

        # ---------------------------------------------------------
        # 4. Create chunks
        # ---------------------------------------------------------
        chunks = []
        next_number = 1

        for page in pages:
            page_chunks = chunk_text(
                page["text"],
                paper_id,
                page["page_number"],
                start_number=next_number,
            )

            chunks.extend(page_chunks)
            next_number += len(page_chunks)

        if not chunks:
            raise ValueError(
                "No text chunks could be created from this file."
            )

        # ---------------------------------------------------------
        # 5. Generate embeddings
        # ---------------------------------------------------------
        embeddings = embed(
            [chunk["text"] for chunk in chunks]
        )

        # ---------------------------------------------------------
        # 6. Save paper metadata
        # ---------------------------------------------------------
        saved = save_paper_if_unique(
            paper_id,
            safe_name,
            str(file_path),
            page_count,
            user["id"],
            content_sha256,
            extension,
        )

        if not saved["created"]:
            file_path.unlink(missing_ok=True)

            record_activity(
                user["id"],
                "duplicate_upload_skipped",
                {
                    "paper_id": saved["paper"]["id"],
                    "filename": safe_name,
                },
            )

            return duplicate_upload_result(
                saved["paper"]
            )

        # ---------------------------------------------------------
        # 7. Save chunks + embeddings
        # ---------------------------------------------------------
        save_chunks(
            chunks,
            embeddings,
        )

        record_activity(
            user["id"],
            "paper_uploaded",
            {
                "paper_id": paper_id,
                "filename": safe_name,
                "file_type": extension,
            },
        )

    except ValueError as exc:
        file_path.unlink(missing_ok=True)

        raise HTTPException(
            status_code=422,
            detail=str(exc),
        ) from exc

    except Exception as exc:
        file_path.unlink(missing_ok=True)

        raise HTTPException(
            status_code=422,
            detail=f"Could not process file: {exc}",
        ) from exc

    return {
        "paper_id": paper_id,
        "filename": safe_name,
        "file_type": extension,
        "total_pages": page_count,
        "total_chunks": len(chunks),
        "status": "ready",
    }

@app.post("/api/chat")
def chat(request: ChatRequest, user: dict[str, Any] = Depends(current_user)):
    if request.paper_ids:
        for paper_id in request.paper_ids:
            owned_paper(paper_id, user)
    hits = hybrid_search(request.question, embed_query(request.question), request.paper_ids, request.top_k, owner_id=user["id"])
    citations = citations_for(hits)
    session = None
    history: list[dict[str, str]] = []
    if request.session_id:
        # Only recent turns are sent back to the model. Fetching the full transcript
        # on every turn gets slower as a chat grows.
        session = get_user_chat(user["id"], request.session_id, message_limit=8)
        if not session:
            raise HTTPException(status_code=404, detail="Chat not found")
        history = [{"role": message["role"], "content": message["content"]}
                   for message in session["messages"][-8:]
                   if message["role"] in {"user", "assistant"}]
    else:
        session = create_chat_session(user["id"], request.question.strip()[:120], request.paper_ids[0] if request.paper_ids else None)
    save_chat_message(session["id"], "user", request.question)
    response = {"answer": answer(request.question, hits, history), "citations": citations,
                "sources": [{"type": "page", "label": f"{item['paperTitle']} p.{item['page']}", "page": item["page"]} for item in citations],
                "mode": "llm" if ai_provider_status()["configured"] else "retrieval-only"}
    save_chat_message(session["id"], "assistant", response["answer"], citations, response["sources"])
    response["session_id"] = session["id"]
    record_activity(user["id"], "asked_question", {"paper_ids": request.paper_ids or []})
    return response


@app.post("/api/learning/generate", status_code=201)
def generate_learning(payload: GenerateRequest, user: dict[str, Any] = Depends(current_user)):
    custom_presentation = payload.kind == "ppt_outline" and payload.presentation_input_type in {"notes", "topic"}
    paper_ids = list(dict.fromkeys(payload.paper_ids or ([payload.paper_id] if payload.paper_id else [])))
    if custom_presentation:
        source_text = payload.source_text.strip()
        if not source_text:
            raise HTTPException(status_code=422, detail="Enter notes or a presentation topic first")
        paper_ids = []
        papers = []
        paper_names = "your provided notes" if payload.presentation_input_type == "notes" else source_text[:100]
        source_label = "User-provided notes" if payload.presentation_input_type == "notes" else "User-provided topic"
        hits = [
            {
                "paper_id": None,
                "filename": source_label,
                "page_number": "n/a",
                "section": f"Input {index + 1}",
                "content": source_text[offset:offset + 1600],
            }
            for index, offset in enumerate(range(0, len(source_text), 1600))
        ]
    else:
        if not paper_ids:
            raise HTTPException(status_code=422, detail="Select an uploaded paper")
        papers = [owned_paper(paper_id, user) for paper_id in paper_ids]
        paper_names = ", ".join(paper["filename"] for paper in papers)
    focus = payload.prompt.strip()
    retrieval_topics = {
            "quiz": "important concepts definitions methods results findings conclusions limitations",
            "flashcards": "key concepts definitions terminology methods findings takeaways",
            "mindmap": "central topic key concepts themes methods results relationships",
            "comparison": "research objectives methodology datasets experiments results limitations contributions",
            "literature_review": "research themes methods findings results limitations trends",
            "research_gap": "limitations unresolved questions future work missing evidence contradictory findings",
            "research_ideas": "limitations unresolved questions future work methods findings",
            "reading_guide": "abstract introduction research question contribution methodology experiments results discussion limitations conclusion",
            "method_extraction": "methodology research design dataset data sample participants baselines metrics evaluation experimental setup",
    }
    topic_query = retrieval_topics.get(payload.kind)
    if custom_presentation:
        query = payload.source_text.strip()
    elif topic_query:
        query = f"{topic_query} {focus} from {paper_names}".strip()
    elif focus:
        query = focus
    else:
        query = f"key findings and main ideas from {paper_names}"
    if not custom_presentation and payload.kind == "ppt_outline":
        # Presentations should represent the full selected source.
        hits = list_paper_chunks(paper_ids, user["id"])
    elif not custom_presentation:
        query_embedding = embed_query(query)
        cross_paper_kinds = {"comparison", "literature_review", "research_gap", "research_ideas"}
        if payload.kind in {"flashcards", "quiz"}:
            focus_terms = f"{focus} {paper_names}".strip()
            retrieval_queries = [
                f"definitions terminology concepts background motivation {focus_terms}",
                f"methods methodology materials procedures design {focus_terms}",
                f"results findings evidence measurements outcomes {focus_terms}",
                f"discussion limitations conclusions implications {focus_terms}",
            ]
            if payload.kind == "quiz" and payload.variation_seed:
                query_offset = sum(ord(character) for character in payload.variation_seed) % len(retrieval_queries)
                retrieval_queries = retrieval_queries[query_offset:] + retrieval_queries[:query_offset]
            per_query_limit = (6 if payload.kind == "flashcards"
                               else min(max((payload.count + 3) // 4, 4), 8))
            hits = []
            seen_chunks = set()
            for retrieval_query in retrieval_queries:
                retrieval_embedding = query_embedding if retrieval_query == query else embed_query(retrieval_query)
                for hit in hybrid_search(retrieval_query, retrieval_embedding, paper_ids, per_query_limit, owner_id=user["id"]):
                    chunk_key = (hit["paper_id"], hit.get("chunk_number"), hit["page_number"])
                    if chunk_key in seen_chunks:
                        continue
                    seen_chunks.add(chunk_key)
                    hits.append(hit)
                    if len(hits) >= 32:
                        break
                if len(hits) >= 32:
                    break
        elif payload.kind in cross_paper_kinds and len(paper_ids) > 1:
            per_paper_limit = max(1, min(4, payload.count // len(paper_ids)))
            hits = []
            for selected_id in paper_ids:
                hits.extend(hybrid_search(query, query_embedding, [selected_id], per_paper_limit, owner_id=user["id"]))
        else:
            retrieval_limit = min(payload.count, 10)
            hits = hybrid_search(query, query_embedding, paper_ids, retrieval_limit, owner_id=user["id"])
    if not hits:
        raise HTTPException(status_code=404, detail="No text passages found for this paper")
    try:
        if payload.kind == "mindmap" and focus:
            stop_words = {"about", "from", "with", "this", "that", "paper", "focus", "map", "mind", "please", "explain", "show", "create", "generate", "the", "and", "for"}
            topic_terms = {term for term in re.findall(r"[a-z0-9]+", focus.casefold())
                           if len(term) > 2 and term not in stop_words}
            evidence_terms = {term for hit in hits for term in re.findall(r"[a-z0-9]+", str(hit.get("content", "")).casefold())}
            if topic_terms and not topic_terms.intersection(evidence_terms):
                raise HTTPException(status_code=422, detail=f'The selected paper does not contain evidence for "{focus}". Try a topic covered by the paper.')
        generation_prompt = focus if payload.kind == "mindmap" else query
        content = generate_study_artifact(
            payload.kind,
            focus or ("Create a balanced, source-wide presentation covering the major sections, methods, findings, and limitations in the provided material."
                      if payload.kind == "ppt_outline" else query),
            hits,
            payload.count,
            payload.difficulty,
            avoid_questions=payload.avoid_questions,
            variation_seed=payload.variation_seed,
            presentation_purpose=payload.presentation_purpose,
            presentation_input_type=payload.presentation_input_type,
        )
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    if content.get("_generation_mode") == "source_fallback" and payload.kind != "ppt_outline":
        raise HTTPException(status_code=503, detail=content.get("_generation_notice", "AI generation failed. Please try again."))
    if payload.kind == "quiz":
        questions = content.get("questions", [])
        # Shuffle on the server before persistence. The browser receives only
        # options; answer keys stay in the database until quiz submission.
        for question in content.get("questions", []):
            options = question.get("options") if isinstance(question, dict) else None
            answer_index = question.get("answer") if isinstance(question, dict) else None
            if isinstance(options, list) and type(answer_index) is int and 0 <= answer_index < len(options):
                correct_option = options[answer_index]
                shuffled_options = list(options)
                random.shuffle(shuffled_options)
                question["options"] = shuffled_options
                question["answer"] = shuffled_options.index(correct_option)
    kind_title = payload.kind.replace("_", " ").title()
    artifact_title = (f"{kind_title}: {focus}" if focus else f"{kind_title} from {paper_names}")[:200]
    if custom_presentation:
        for item in content.get("slides", []):
            item["page"] = None
    elif payload.kind == "ppt_outline":
        valid_pages = {int(hit["page_number"]) for hit in hits if isinstance(hit.get("page_number"), int)}
        for item in content.get("slides", []):
            page = item.get("page")
            if not isinstance(page, int) or page not in valid_pages:
                item["page"] = None
    citations = [] if custom_presentation or payload.kind == "ppt_outline" else citations_for(hits)
    artifact_payload = {**content, "citations": citations, "source_paper_ids": paper_ids}
    if payload.kind == "ppt_outline":
        artifact_payload.update({
            "presentation_input_type": payload.presentation_input_type,
            "presentation_purpose": payload.presentation_purpose,
        })
    artifact = save_artifact(user["id"], payload.kind, artifact_title, artifact_payload, paper_ids[0] if paper_ids else None)
    record_activity(user["id"], f"generated_{payload.kind}", {"artifact_id": artifact["id"], "paper_ids": paper_ids})
    if payload.kind == "quiz":
        public_payload = dict(artifact["payload"])
        public_payload["questions"] = [
            {key: value for key, value in question.items() if key != "answer"}
            for question in public_payload.get("questions", [])
        ]
        artifact = {**artifact, "payload": public_payload}
    return {"artifact": artifact}


@app.post("/api/artifacts/{artifact_id}/flashcard-reviews")
def save_flashcard_reviews(artifact_id: str, payload: FlashcardReviewRequest,
                           user: dict[str, Any] = Depends(current_user)):
    artifact = get_user_artifact(user["id"], artifact_id)
    if not artifact or artifact["kind"] != "flashcards":
        raise HTTPException(status_code=404, detail="Flashcard deck not found")
    cards = artifact.get("payload", {}).get("cards", [])
    for index, rating in payload.reviews.items():
        if not index.isdecimal() or int(index) >= len(cards):
            raise HTTPException(status_code=422, detail="A reviewed flashcard was not found in this deck")
        if rating not in {"again", "hard", "got-it"}:
            raise HTTPException(status_code=422, detail="Flashcard rating must be Again, Hard, or Got It")
    if not update_flashcard_reviews(user["id"], artifact_id, payload.reviews):
        raise HTTPException(status_code=404, detail="Flashcard deck not found")
    record_activity(user["id"], "reviewed_flashcard", {"artifact_id": artifact_id, "reviews": len(payload.reviews)})
    return {"ok": True}


@app.post("/api/quizzes/{artifact_id}/attempts", status_code=201)
def submit_quiz_attempt(artifact_id: str, payload: QuizAttemptRequest,
                        user: dict[str, Any] = Depends(current_user)):
    artifact = get_user_artifact(user["id"], artifact_id)
    if not artifact or artifact["kind"] != "quiz":
        raise HTTPException(status_code=404, detail="Quiz not found")
    questions = artifact.get("payload", {}).get("questions", [])
    if not isinstance(questions, list) or not questions:
        raise HTTPException(status_code=422, detail="This quiz has no questions to submit")
    correct = 0
    validated_answers: dict[str, str] = {}
    for question_index, question in enumerate(questions):
        options = question.get("options", [])
        answer_index = question.get("answer")
        if not isinstance(options, list) or type(answer_index) is not int or not 0 <= answer_index < len(options):
            raise HTTPException(status_code=422, detail="This quiz contains an invalid answer key")
        selected = payload.answers.get(str(question_index))
        if selected is not None and selected not in options:
            raise HTTPException(status_code=422, detail="Choose one of the listed answer options")
    for question_index, selected in payload.answers.items():
        if not question_index.isdecimal() or int(question_index) >= len(questions):
            raise HTTPException(status_code=422, detail="An answer refers to a question that is not in this quiz")
        question = questions[int(question_index)]
        options = question.get("options", [])
        if not isinstance(selected, str) or selected not in options:
            raise HTTPException(status_code=422, detail="Choose one of the listed answer options")
        answer_index = question.get("answer")
        if not isinstance(answer_index, int) or not 0 <= answer_index < len(options):
            raise HTTPException(status_code=422, detail="This quiz contains an invalid answer key")
        validated_answers[question_index] = selected
        if selected == options[answer_index]:
            correct += 1
    total = len(questions)
    unanswered = total - len(validated_answers)
    incorrect = len(validated_answers) - correct
    attempt = save_quiz_attempt(user["id"], artifact_id, total, correct, incorrect, unanswered, validated_answers)
    if not attempt:
        raise HTTPException(status_code=404, detail="Quiz not found")
    attempt_answers = attempt.pop("answers", {})
    created = attempt.pop("created", False)
    question_results = []
    for question_index, question in enumerate(questions):
        options = question["options"]
        answer_index = question["answer"]
        correct_option = options[answer_index]
        selected = attempt_answers.get(str(question_index))
        question_results.append({
            "question_index": question_index,
            "selected_answer": selected,
            "correct_answer": correct_option,
            "is_correct": selected == correct_option,
            "explanation": question.get("explanation", ""),
        })
    if created:
        record_activity(user["id"], "submitted_quiz", {"artifact_id": artifact_id, "attempt_id": attempt["id"], "percentage": float(attempt["percentage"])})
    return {"attempt": {**attempt, "questions": question_results}}


@app.post("/api/learning/mindmap/expand", response_model=MindMapExpansionResponse)
def expand_learning_mindmap(payload: MindMapExpansionRequest, user: dict[str, Any] = Depends(current_user)):
    owned_paper(payload.paper_id, user)
    existing_nodes = [node.model_dump() for node in payload.existing_nodes]
    if not any(node["id"] == payload.node_id for node in existing_nodes):
        raise HTTPException(status_code=422, detail="The selected concept is not part of this mind map")

    retrieval_query = " ".join([
        payload.node_label,
        *payload.breadcrumb[-6:],
        payload.context[:700],
        "definition mechanism key details examples advantages limitations applications relationships",
    ])
    hits = hybrid_search(retrieval_query, embed_query(retrieval_query), [payload.paper_id], 12, owner_id=user["id"])
    if not hits:
        raise HTTPException(status_code=404, detail="No paper passages found for this concept")
    try:
        return expand_mindmap_node(
            payload.node_id, payload.node_label, payload.breadcrumb, payload.context, existing_nodes, hits,
        )
    except ValueError as exc:
        raise HTTPException(status_code=503, detail=str(exc)) from exc


@app.post("/api/presentations/export")
def export_presentation(payload: PresentationExportRequest, user: dict[str, Any] = Depends(current_user)):
    """Build a PowerPoint from the signed-in user's current presentation editor state."""
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)

    for item in payload.slides:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        background = slide.background.fill
        background.solid()
        background.fore_color.rgb = RGBColor(248, 250, 252)

        band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, presentation.slide_width, Inches(0.16))
        band.fill.solid()
        band.fill.fore_color.rgb = RGBColor(79, 70, 229)
        band.line.fill.background()

        title_box = slide.shapes.add_textbox(Inches(0.75), Inches(0.55), Inches(11.8), Inches(1.0))
        title_frame = title_box.text_frame
        title_frame.clear()
        title_frame.word_wrap = True
        title_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        title_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        title_paragraph = title_frame.paragraphs[0]
        title_paragraph.text = item.title.strip()
        title_paragraph.font.name = "Aptos Display"
        title_paragraph.font.size = Pt(30)
        title_paragraph.font.bold = True
        title_paragraph.font.color.rgb = RGBColor(15, 23, 42)

        content_box = slide.shapes.add_textbox(Inches(1.0), Inches(1.9), Inches(11.2), Inches(4.75))
        content_frame = content_box.text_frame
        content_frame.clear()
        content_frame.word_wrap = True
        content_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        lines = [line.strip() for line in item.content if line.strip()]
        font_size = 20 if len(lines) <= 6 else 17
        if not lines:
            lines = ["No slide content provided."]
        for line_index, line in enumerate(lines):
            paragraph = content_frame.paragraphs[0] if line_index == 0 else content_frame.add_paragraph()
            paragraph.text = f"•  {line}"
            paragraph.level = 0
            paragraph.font.name = "Aptos"
            paragraph.font.size = Pt(font_size)
            paragraph.font.color.rgb = RGBColor(51, 65, 85)
            paragraph.space_after = Pt(14)

        footer = slide.shapes.add_textbox(Inches(0.75), Inches(7.02), Inches(11.8), Inches(0.25))
        footer_paragraph = footer.text_frame.paragraphs[0]
        cited_pages = sorted(set(item.pages + ([item.page] if item.page else [])))
        source = f"Source: pages {', '.join(map(str, cited_pages))}" if cited_pages else payload.source_label
        footer_paragraph.text = f"ScholarMind  •  {source}"
        footer_paragraph.font.name = "Aptos"
        footer_paragraph.font.size = Pt(10)
        footer_paragraph.font.color.rgb = RGBColor(100, 116, 139)
        footer_paragraph.alignment = PP_ALIGN.LEFT

        notes = item.speaker_notes.strip()
        if cited_pages:
            notes = f"Source pages: {', '.join(map(str, cited_pages))}\n\n{notes}".strip()
        if notes:
            slide.notes_slide.notes_text_frame.text = notes

    output = BytesIO()
    presentation.save(output)
    output.seek(0)
    filename = re.sub(r"[^A-Za-z0-9_-]+", "-", payload.title).strip("-_").lower() or "scholarmind-presentation"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        headers={"Content-Disposition": f'attachment; filename="{filename[:80]}.pptx"'},
    )


@app.patch("/api/artifacts/{artifact_id}")
def update_artifact(artifact_id: str, payload: ArtifactUpdate, user: dict[str, Any] = Depends(current_user)):
    from database import update_user_artifact
    result = update_user_artifact(user["id"], artifact_id, payload.title, payload.payload)
    if not result:
        raise HTTPException(status_code=404, detail="Saved item not found")
    record_activity(user["id"], "updated_saved_item", {"artifact_id": artifact_id})
    return {"artifact": result}


@app.delete("/api/artifacts/{artifact_id}", status_code=204)
def delete_artifact(artifact_id: str, user: dict[str, Any] = Depends(current_user)):
    from database import delete_user_artifact
    if not delete_user_artifact(user["id"], artifact_id):
        raise HTTPException(status_code=404, detail="Saved item not found")
    record_activity(user["id"], "deleted_saved_item", {"artifact_id": artifact_id})
