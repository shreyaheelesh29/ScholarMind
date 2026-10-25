from __future__ import annotations

import os
import shutil
from io import BytesIO
from pathlib import Path
from typing import Any

import pymupdf
import pytesseract
from PIL import Image
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


if os.name == "nt" and not shutil.which("tesseract"):
    tesseract_path = Path(os.environ.get("ProgramFiles", r"C:\Program Files")) / "Tesseract-OCR" / "tesseract.exe"
    if tesseract_path.is_file():
        pytesseract.pytesseract.tesseract_cmd = str(tesseract_path)


SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".pptx",
    ".txt",
    ".png",
    ".jpg",
    ".jpeg",
}


def get_file_extension(filename: str) -> str:
    return Path(filename).suffix.lower()


def is_supported_file(filename: str) -> bool:
    return get_file_extension(filename) in SUPPORTED_EXTENSIONS


def extract_pdf(file_path: Path) -> dict[str, Any]:
    pages = []

    with pymupdf.open(file_path) as document:
        for page_number, page in enumerate(document, start=1):
            text = page.get_text("text").strip()
            # OCR fallback for scanned/image-only PDF pages
            if not text:
                pix = page.get_pixmap()
                image = Image.frombytes(
                    "RGB",
                    [pix.width, pix.height],
                    pix.samples,
                )

                text = pytesseract.image_to_string(image).strip()

            pages.append({"page_number": page_number, "text": text})

        page_count = len(document)

    return {
        "text": "\n\n".join(page["text"] for page in pages),
        "pages": pages,
        "page_count": page_count,
        "document_type": "pdf",
    }


def extract_docx(file_path: Path) -> dict[str, Any]:
    document = Document(file_path)
    paragraphs = []
    # Walk the document body in order so tables stay beside the section that
    # explains them, and retain heading cues for downstream chunking/retrieval.
    for element in document.element.body.iterchildren():
        if element.tag.endswith("}p"):
            paragraph = Paragraph(element, document)
            text = paragraph.text.strip()
            if not text:
                continue
            style_name = getattr(paragraph.style, "name", "") or ""
            if style_name.lower().startswith("heading"):
                paragraphs.append(f"\n{style_name}: {text}")
            else:
                paragraphs.append(text)
        elif element.tag.endswith("}tbl"):
            table = Table(element, document)
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells)
                if row_text.strip(" |"):
                    paragraphs.append(row_text)

    full_text = "\n\n".join(paragraphs)

    return {
        "text": full_text,
        "pages": [
            {
                "page_number": 1,
                "text": full_text,
            }
        ]
        if full_text
        else [],
        "page_count": 1 if full_text else 0,
        "document_type": "docx",
    }


def extract_pptx(file_path: Path) -> dict[str, Any]:
    presentation = Presentation(file_path)

    slides = []

    def collect_shape_text(shapes) -> list[str]:
        collected = []
        for shape in shapes:
            # Grouped objects and tables often hold slide text that is not
            # exposed through the ordinary shape.text property.
            nested_shapes = getattr(shape, "shapes", None)
            if nested_shapes is not None:
                collected.extend(collect_shape_text(nested_shapes))
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    for cell in row.cells:
                        value = cell.text.strip()
                        if value:
                            collected.append(value)
            elif hasattr(shape, "text"):
                value = shape.text.strip()
                if value:
                    collected.append(value)
        return collected

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1,
    ):
        slide_text = collect_shape_text(slide.shapes)

        # Slides that are screenshots or scanned images have no shape text.
        # OCR their embedded pictures so they can contribute to the viewer and
        # source-grounded summaries when Tesseract is available.
        if not slide_text:
            for shape in slide.shapes:
                if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
                    continue
                try:
                    image = Image.open(BytesIO(shape.image.blob)).convert("RGB")
                    image.thumbnail((2400, 2400))
                    recognized = pytesseract.image_to_string(image).strip()
                    if recognized:
                        slide_text.append(recognized)
                except Exception:
                    # Some embedded formats are not readable by Pillow; keep
                    # the slide in the index even if it has no extractable text.
                    continue

        combined_text = "\n".join(slide_text).strip()

        # Preserve physical slide numbering, including image-only slides. The
        # viewer and citations use page_number as a 1-based slide index.
        slides.append({"page_number": slide_number, "text": combined_text})

    return {
        "text": "\n\n".join(
            slide["text"] for slide in slides
        ),
        "pages": slides,
        "page_count": len(presentation.slides),
        "document_type": "pptx",
    }


def extract_txt(file_path: Path) -> dict[str, Any]:
    text = file_path.read_text(
        encoding="utf-8",
        errors="replace",
    ).strip()

    return {
        "text": text,
        "pages": [
            {
                "page_number": 1,
                "text": text,
            }
        ]
        if text
        else [],
        "page_count": 1 if text else 0,
        "document_type": "txt",
    }


def extract_image(file_path: Path) -> dict[str, Any]:
    image = Image.open(file_path)

    # Convert image to grayscale
    image = image.convert("L")

    # Simple thresholding for better OCR
    image = image.point(
        lambda x: 0 if x < 150 else 255,
        "1",
    )

    text = pytesseract.image_to_string(image).strip()

    return {
        "text": text,
        "pages": [
            {
                "page_number": 1,
                "text": text,
            }
        ]
        if text
        else [],
        "page_count": 1,
        "document_type": "image",
    }


def extract_text(file_path: Path) -> dict[str, Any]:
    extension = get_file_extension(file_path.name)

    if extension == ".pdf":
        return extract_pdf(file_path)

    if extension == ".docx":
        return extract_docx(file_path)

    if extension == ".pptx":
        return extract_pptx(file_path)

    if extension == ".txt":
        return extract_txt(file_path)

    if extension in {".png", ".jpg", ".jpeg"}:
        return extract_image(file_path)

    raise ValueError(
        f"Unsupported file type: {extension}"
    )
