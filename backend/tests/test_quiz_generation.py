from __future__ import annotations

import json
import re
import sys
import tempfile
import threading
import unittest
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import pymupdf
from docx import Document

import document_processor
import chunking
import rag


def make_hits(count: int = 32) -> list[dict]:
    return [
        {"paper_id": "paper", "filename": "study.pdf", "page_number": index + 1,
         "section": f"Section {index + 1}",
         "content": f"Concept {index} describes a distinct method, result, and implication in detail."}
        for index in range(count)
    ]


def make_question(number: int, page: int = 1) -> dict:
    return {
        "question": f"How does concept {number} affect the reported study outcome?",
        "options": [f"Outcome {number}A", f"Outcome {number}B", f"Outcome {number}C", f"Outcome {number}D"],
        "answer": 0,
        "explanation": f"The source reports this effect for concept {number}.",
        "page": page,
    }


class QuizGenerationTests(unittest.TestCase):
    def generate_with_mock(self, count: int, mock_request, local: bool = False):
        settings = ("", "http://127.0.0.1:11434", True, "test-model") if local else (
            "test-key", "https://example.invalid/v1", False, "test-model")
        with patch.object(rag, "_provider_settings", return_value=settings), \
             patch.object(rag, "_chat_request", side_effect=mock_request), \
             patch.object(rag, "_questions_are_equivalent", side_effect=lambda a, b: rag._question_key(a) == rag._question_key(b)):
            return rag.generate_study_artifact("quiz", "study concepts", make_hits(), count)

    def test_requested_counts_of_five_ten_and_twenty_are_dynamic(self):
        for count in (5, 10, 20):
            with self.subTest(count=count):
                next_number = 0

                def model(_base_url, _model, messages, **kwargs):
                    nonlocal next_number
                    text = messages[-1]["content"]
                    additional = re.search(r"Create exactly (\d+) additional", text)
                    initial = re.search(r"Generate exactly (\d+) distinct questions", text)
                    amount = int(additional.group(1)) if additional else int(initial.group(1))
                    if not additional and count > 5:
                        amount = min(5, amount)
                    questions = []
                    for _ in range(amount):
                        next_number += 1
                        questions.append(make_question(next_number))
                    return json.dumps({"questions": questions})

                result = self.generate_with_mock(count, model)
                self.assertEqual(len(result["questions"]), count)

    def test_partial_initial_batch_requests_only_missing_questions(self):
        calls = []
        next_number = 0

        def model(_base_url, _model, messages, **kwargs):
            nonlocal next_number
            text = messages[-1]["content"]
            calls.append(text)
            additional = re.search(r"Create exactly (\d+) additional", text)
            initial = re.search(r"Generate exactly (\d+) distinct questions", text)
            amount = int(additional.group(1)) if additional else min(5, int(initial.group(1)))
            questions = []
            for _ in range(amount):
                next_number += 1
                questions.append(make_question(next_number))
            return json.dumps({"questions": questions})

        result = self.generate_with_mock(10, model)
        self.assertEqual(len(result["questions"]), 10)
        self.assertEqual(len(calls), 3)  # one initial batch and two missing-only batches
        self.assertIn("Create exactly 4 additional", calls[1])

    def test_local_partial_batches_do_not_run_nested_repair_loops(self):
        calls = []
        next_number = 0

        def model(_base_url, _model, messages, **kwargs):
            nonlocal next_number
            text = messages[-1]["content"]
            calls.append(text)
            requested = re.search(r"Generate exactly (\d+) distinct questions", text)
            amount = min(2, int(requested.group(1)))
            questions = []
            for _ in range(amount):
                next_number += 1
                questions.append(make_question(next_number))
            return json.dumps({"questions": questions})

        result = self.generate_with_mock(10, model, local=True)
        self.assertEqual(len(result["questions"]), 10)
        self.assertEqual(len(calls), 5)
        self.assertTrue(all("Create exactly" not in prompt for prompt in calls))

    def test_duplicate_items_are_removed_but_distinct_aspects_remain(self):
        first = make_question(1)
        duplicate = dict(first, question=first["question"])
        different_aspect = make_question(2)
        with patch.object(rag, "_provider_settings", return_value=("key", "https://example.invalid/v1", False, "test")), \
             patch.object(rag, "_chat_request", return_value=json.dumps({"questions": [first, duplicate, different_aspect]})):
            result = rag.generate_study_artifact("quiz", "study", make_hits(), 2)
        self.assertEqual(len(result["questions"]), 2)

    def test_insufficient_material_returns_partial_after_bounded_batches(self):
        calls = 0
        first = make_question(1)
        second = make_question(2)

        def model(_base_url, _model, messages, **kwargs):
            nonlocal calls
            calls += 1
            return json.dumps({"questions": [first, second]})

        result = self.generate_with_mock(10, model, local=True)
        self.assertEqual(len(result["questions"]), 2)
        self.assertIn("of 10", result["_generation_notice"])
        self.assertLessEqual(calls, 12)

    def test_provider_timeout_returns_controlled_fallback(self):
        def timeout(*args, **kwargs):
            raise TimeoutError("mock timeout")

        result = self.generate_with_mock(5, timeout)
        self.assertEqual(result["_generation_mode"], "source_fallback")
        self.assertIn("timed out", result["_generation_notice"])

    def test_concurrent_requests_keep_independent_question_sets(self):
        sequence = iter(range(1, 1000))
        lock = threading.Lock()

        def model(_base_url, _model, messages, **kwargs):
            match = re.search(r"Generate exactly (\d+) distinct questions", messages[-1]["content"])
            amount = int(match.group(1))
            with lock:
                questions = [make_question(next(sequence)) for _ in range(amount)]
            return json.dumps({"questions": questions})

        with patch.object(rag, "_provider_settings", return_value=("key", "https://example.invalid/v1", False, "test")), \
             patch.object(rag, "_chat_request", side_effect=model):
            with ThreadPoolExecutor(max_workers=2) as pool:
                results = list(pool.map(lambda _: rag.generate_study_artifact("quiz", "study", make_hits(), 5), range(2)))
        self.assertEqual([len(result["questions"]) for result in results], [5, 5])
        for result in results:
            self.assertEqual(len({item["question"] for item in result["questions"]}), 5)


class DocumentExtractionTests(unittest.TestCase):
    def test_docx_preserves_heading_and_table_order(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "paper.docx"
            document = Document()
            document.add_heading("Abstract", level=1)
            document.add_paragraph("Abstract content.")
            table = document.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "Metric"
            table.cell(0, 1).text = "Value"
            document.add_heading("Results", level=1)
            document.add_paragraph("Results content.")
            document.save(path)

            extracted = document_processor.extract_text(path)
            self.assertLess(extracted["text"].index("Heading 1: Abstract"), extracted["text"].index("Metric | Value"))
            self.assertLess(extracted["text"].index("Metric | Value"), extracted["text"].index("Heading 1: Results"))
            chunks = chunking.chunk_text(extracted["text"], "paper", 1, chunk_size=200)
            self.assertIn("Abstract", {chunk["section"] for chunk in chunks})
            self.assertIn("Results", {chunk["section"] for chunk in chunks})

    def test_pdf_extraction_keeps_page_text(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "paper.pdf"
            document = pymupdf.open()
            page = document.new_page()
            page.insert_text((72, 72), "A paper-supported result on page one.")
            document.save(path)
            document.close()

            extracted = document_processor.extract_text(path)
            self.assertEqual(extracted["page_count"], 1)
            self.assertIn("paper-supported result", extracted["pages"][0]["text"])


if __name__ == "__main__":
    unittest.main()
